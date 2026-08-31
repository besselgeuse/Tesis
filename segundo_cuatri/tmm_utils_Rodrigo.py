
# -*- coding: utf-8 -*-
"""
Created on Thu Mar 26 17:45:17 2020

@author: simon

Utilities to calculate RAT of optical stacks with tmm_vec
"""
# 8/4 cambie los imports de Simon para que se pueda elegir la ruta del tmm_core

import numpy as np
import os
import re
import ast
import pickle
import sys
import importlib.util
import torch
import torch.optim as optim

# 1. Definimos la ruta a tmm_core.py (priorizando el mismo directorio)
import os
_dir = os.path.dirname(__file__)
tmm_path = os.path.join(_dir, "tmm_core.py")
if not os.path.exists(tmm_path):
    tmm_path = os.path.abspath(os.path.join(_dir, "..", "tmm_core.py"))

# 2. Cargamos el módulo manualmente desde esa dirección
spec = importlib.util.spec_from_file_location("tmm", tmm_path)
tmm = importlib.util.module_from_spec(spec)
sys.modules["tmm"] = tmm # Esto asegura que otros scripts vean esta versión
spec.loader.exec_module(tmm)

from functools import partial
from itertools import product
from scipy.interpolate import interp1d
from openpyxl import load_workbook
import matplotlib.pyplot as plt


def cauchy_fn(A,B,C):
    def fn(lams):
        lams = lams*1.0
        return A+B*np.power(10,4)/np.power(lams,2)+C*np.power(10,9)/np.power(lams,4)
    return fn


def constant_fn(val):
    return partial(np.full_like, fill_value=val)

def n_eff(n1,n2,f):
    """
    Modelo de bruggeman. f es la fraccion de la especie 2.
    """
    if f == 0: return n1
    if f == 1: return n2
    e1 = n1**2
    e2 = n2**2
    #f = 1-f2
    omega = (1-f)*(e2-2*e1)+f*(e1-2*e2)
    return (np.sqrt(np.sqrt(omega**2 + 8*e1*e2)-omega))/2


def brugg_fn(n_a, n_b, f_b):
    def fn(lams):
        n1 = n_a(lams)
        n2 = n_b(lams)
        return n_eff(n1, n2, f_b)
    return fn

def n_eff_torch(n1, n2, f):
    e1 = n1**2
    e2 = n2**2
    # Ecuación de Bruggeman simétrica
    omega = (1-f)*(e2 - 2*e1) + f*(e1 - 2*e2)
    # sqrt de PyTorch sobre tensores complejos es nativa
    return torch.sqrt((torch.sqrt(omega**2 + 8*e1*e2) - omega) / 4.0)

def brugg_fn_torch(n_a, n_b, f_b):
    def fn(lams):
        # Asegúrate de que n_a y n_b devuelvan tensores PyTorch
        n1 = n_a(lams)
        n2 = n_b(lams)
        return n_eff_torch(n1, n2, f_b)
    return fn

def load_fn(filename, skiprows=1,delimiter = "\t"):
    
    wv, I = np.loadtxt(filename, skiprows=skiprows,delimiter=f"{delimiter}" ,unpack=True)
    return interp1d(wv, I,fill_value=(I[0],I[-1]), bounds_error=False)


def load_interp(filename, comments=';', skiprows=0, unit='nm'):
    '''loads optical data and outputs interpolation functions for n and k. 
    Output functions take wv in nm. Keep default arguments to load .nkv files
    '''
    scale = {'nm':1.0, 'um':1000.0}
    
    wv, n, k = np.loadtxt(filename, comments=comments, 
                      skiprows=skiprows, unpack=True)
    
    wv *= scale[unit]
    
    n_fn = interp1d(wv, n,fill_value=(n[0],n[-1]), bounds_error=False)
    k_fn = interp1d(wv, k,fill_value=(k[0],k[-1]), bounds_error=False)
    
    return n_fn, k_fn


def load_interp_in3(filename, unit='A'):
    '''loads optical data and outputs interpolation functions for n and k from .in3 files.
    Output functions take wv in nm. Default unit is Angstroms ('A').
    '''
    scale = {'nm': 1.0, 'um': 1000.0, 'A': 0.1}
    
    with open(filename, 'r') as f:
        lines = [line.strip() for line in f.readlines() if line.strip() != '']
    
    n_pts = 0
    k_pts = 0
    
    wv_n = []
    n_vals = []
    wv_k = []
    k_vals = []
    
    idx = 0
    # Buscar el número de puntos de n
    while idx < len(lines):
        parts = lines[idx].split()
        if len(parts) == 1:
            try:
                n_pts = int(parts[0])
                idx += 1
                break
            except ValueError:
                pass
        idx += 1
        
    for _ in range(n_pts):
        parts = lines[idx].split()
        wv_n.append(float(parts[0]))
        n_vals.append(float(parts[1]))
        idx += 1
        
    # Buscar el número de puntos de k
    while idx < len(lines):
        parts = lines[idx].split()
        if len(parts) == 1:
            try:
                k_pts = int(parts[0])
                idx += 1
                break
            except ValueError:
                pass
        idx += 1
        
    for _ in range(k_pts):
        parts = lines[idx].split()
        wv_k.append(float(parts[0]))
        k_vals.append(float(parts[1]))
        idx += 1
        
    wv_n = np.array(wv_n)
    wv_k = np.array(wv_k)
    n_vals = np.array(n_vals)
    k_vals = np.array(k_vals)
    
    if unit in scale:
        wv_n *= scale[unit]
        wv_k *= scale[unit]
        
    n_fn = interp1d(wv_n, n_vals, fill_value=(n_vals[0], n_vals[-1]), bounds_error=False)
    k_fn = interp1d(wv_k, k_vals, fill_value=(k_vals[0], k_vals[-1]), bounds_error=False)
    
    return n_fn, k_fn


def guardar_resultados_txt(salida_path, data_file, wl_exp, Is_exp, Is_fit, Ic_exp, Ic_fit, 
                           psi_exp, psi_fit, delta_exp, delta_fit, 
                           best_thicknesses, best_params):
    """Guarda los resultados del ajuste en un archivo TXT."""
    with open(salida_path, 'w', encoding='utf-8') as f:
        f.write("# ========================================================\n")
        f.write("# RESULTADOS DE AJUSTE ELIPSOMÉTRICO (AUTORANGE + TMM)\n")
        f.write("# ========================================================\n")
        f.write(f"# Archivo de origen: {os.path.basename(data_file)}\n")
        f.write(f"# Espesores óptimos (nm): {list(best_thicknesses)}\n")
        if best_params:
            f.write(f"# Error mínimo (Chi2 Red / MSE): {best_params.get('chi2_min', 0.0):.6e}\n")
            f.write("# Parámetros de dispersión optimizados:\n")
            for name, params in best_params.items():
                if name != 'chi2_min' and not isinstance(name, int):
                    f.write(f"#   Capa '{name}': {params}\n")
        f.write("# ========================================================\n")
        f.write("# Wavelength_nm\tIs_exp\tIs_fit\tIc_exp\tIc_fit\tPsi_exp_deg\tPsi_fit_deg\tDelta_exp_deg\tDelta_fit_deg\n")
        
        for i in range(len(wl_exp)):
            f.write(f"{wl_exp[i]:.4f}\t{Is_exp[i]:.6f}\t{Is_fit[i]:.6f}\t"
                    f"{Ic_exp[i]:.6f}\t{Ic_fit[i]:.6f}\t{psi_exp[i]:.4f}\t"
                    f"{psi_fit[i]:.4f}\t{delta_exp[i]:.4f}\t{delta_fit[i]:.4f}\n")
    print(f"--> Resultados guardados exitosamente en: {salida_path}")


def leer_datos_guardados(path):
    """
    Carga y reconstruye automáticamente los datos y metadatos guardados en un archivo TXT de resultados.
    
    Parámetros:
    -----------
    path : str
        Ruta al archivo TXT de resultados de ajuste (generado por `guardar_resultados_txt`).
        
    Retorna:
    --------
    data_file : str
        Nombre del archivo original de mediciones elipsométricas.
    wl_exp : np.ndarray
        Vector 1D de longitudes de onda [nm].
    Is_exp : np.ndarray
        Vector 1D de la componente Is experimental.
    best_Is_curve : np.ndarray
        Vector 1D de la componente Is ajustada / calculada (Is_fit).
    Ic_exp : np.ndarray
        Vector 1D de la componente Ic experimental.
    best_Ic_curve : np.ndarray
        Vector 1D de la componente Ic ajustada / calculada (Ic_fit).
    psi_exp : np.ndarray
        Vector 1D del ángulo Psi experimental [grados].
    psi_fit : np.ndarray
        Vector 1D del ángulo Psi ajustado [grados].
    delta_exp : np.ndarray
        Vector 1D del ángulo Delta experimental [grados].
    delta_fit : np.ndarray
        Vector 1D del ángulo Delta ajustado [grados].
    best_thicknesses : np.ndarray
        Vector 1D con los espesores óptimos encontrados [nm].
    best_params : dict
        Diccionario con los parámetros óptimos del ajuste (chi2_min y parámetros por capa).
    """
    data_file = ""
    best_thicknesses = np.array([])
    best_params = {}
    
    with open(path, 'r', encoding='utf-8', errors='replace') as f:
        for line in f:
            line_str = line.strip()
            if not line_str.startswith('#'):
                break
            
            # 1. Archivo de origen
            if 'Archivo de origen:' in line_str:
                data_file = line_str.split('Archivo de origen:', 1)[1].strip()
                
            # 2. Espesores óptimos
            elif 'Espesores' in line_str and ':' in line_str:
                th_str = line_str.split(':', 1)[1].strip()
                try:
                    best_thicknesses = np.array(
                        eval(th_str, {"__builtins__": None, "np": np, "float64": np.float64, "list": list, "tuple": tuple}),
                        dtype=float
                    )
                except Exception:
                    cleaned = re.sub(r'(?:np\.)?float\d*\(|\btensor\(', '', th_str).replace(')', '')
                    matches = re.findall(r'[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?', cleaned)
                    best_thicknesses = np.array([float(x) for x in matches], dtype=float)
                    
            # 3. Error mínimo (Chi2 Red / MSE)
            elif ('Error m' in line_str or 'Error min' in line_str) and ':' in line_str:
                val_str = line_str.split(':', 1)[1].strip()
                try:
                    best_params['chi2_min'] = float(val_str)
                except Exception:
                    pass
                    
            # 4. Parámetros por capa
            elif 'Capa ' in line_str and ':' in line_str:
                m = re.search(r"Capa\s+['\"](.*?)['\"]\s*:\s*(.*)", line_str)
                if m:
                    capa_name = m.group(1).strip()
                    dict_str = m.group(2).strip()
                    try:
                        parsed_dict = eval(
                            dict_str,
                            {"__builtins__": None, "True": True, "False": False, "None": None, "np": np, "float64": np.float64}
                        )
                    except Exception:
                        parsed_dict = ast.literal_eval(dict_str)
                    best_params[capa_name] = parsed_dict

    # Carga de columnas de datos numéricos
    wl_exp, Is_exp, best_Is_curve, Ic_exp, best_Ic_curve, psi_exp, psi_fit, delta_exp, delta_fit = np.loadtxt(
        path, comments='#', unpack=True
    )
    
    return (
        data_file, wl_exp, Is_exp, best_Is_curve, Ic_exp, best_Ic_curve,
        psi_exp, psi_fit, delta_exp, delta_fit,
        best_thicknesses, best_params
    )

def interpolar_todo(wl_exp=None,psi_exp=None,delta_exp=None,Ic_exp=None,Is_exp=None):
    """"
    Esta función tiene el proposito de interpolar todos los datos experimentales 
    utiles medidos por el elipsometro HORIBA de CNEA, en el rango de longitudes de onda
    que van desde los 440 nm hasta los 830 nm, que es el rango de medición valido
    para este elipsometro.
    """"
    def interpolar_datos(lams, datos=None):
        if datos is not None:
            return interp1d(lams,datos,kind='linear',fill_value='extrapolate')
        else:
            return None
    mask_wl = np.where((wl_exp >= 440) & (wl_exp <= 830))
    wl_exp = wl_exp[mask_wl]
    psi_exp = psi_exp[mask_wl]
    delta_exp = delta_exp[mask_wl]
    Ic_exp = Ic_exp[mask_wl]
    Is_exp = Is_exp[mask_wl]

    psi_exp_interp = interpolar_datos(wl_exp,psi_exp)
    delta_exp_interp = interpolar_datos(wl_exp,delta_exp)
    Ic_exp_interp = interpolar_datos(wl_exp,Ic_exp)
    Is_exp_interp = interpolar_datos(wl_exp,Is_exp)

    return wl_exp, psi_exp_interp(wl_exp), delta_exp_interp(wl_exp), Ic_exp_interp(wl_exp), Is_exp_interp(wl_exp)    



def stack2tmm(stack, materials, lams, add_inf=True):
    '''Transforms a stack into three lists that can be fed into the 'tmm' package.
    BY DEFAULT ADDS inf AIR LAYERS'''
    n_list = []
    d_list = []
    c_list = []
    if add_inf:
        n_list.append([1.]*len(lams))
        d_list.append(np.inf)
        c_list.append('i')
        
    for layer in stack:
        d_list.append(layer[0])
        mat = materials[layer[1]]
        n_list.append(mat[0](lams) + mat[1](lams)*1.0j)
        c_list.append(layer[2])
        
    if add_inf:  
        n_list.append([1.]*len(lams))
        d_list.append(np.inf)
        c_list.append('i')
    
    return n_list, d_list, c_list

   
def RT_with_cache(file, stack, materials, lams, pol='s', th_0=0, thicks=None):
    if file is None:
        return calculate_RT(stack, materials, lams, pol, th_0, thicks)
    else:
    
        if os.path.exists(file):
            with open(file,'rb') as f:
                print('loading from cache\n',)
                RT = pickle.load(f)
        else:
            RT = calculate_RT(stack, materials, lams, pol, th_0, thicks)
            with open(file,'wb') as f:
                print('saving to cache\n')
                pickle.dump(RT, f)
        return RT


def calculate_RT(stack, materials, lams, pol='s', th_0=0, thicks=None):
    '''calulate R and T of a stack, for all combinations of thicks.'''
    n_list, d_list, c_list = stack2tmm(stack, materials, lams, add_inf=False)
    n_lams = len(lams)
    inc = 'i' in c_list[1:-1] #check if  any of the finite layers is incocherent
    if inc:
        def f(pol, n_list, d_list, th_0, lams, c_list):
            return tmm.inc_tmm(pol, n_list, d_list, c_list, th_0, lams)
    else:
        def f(pol, n_list, d_list, th_0, lams, c_list=None):
            return tmm.coh_tmm(pol, n_list, d_list, th_0, lams)
            
    if thicks is None:
        #case with only one thickness
        RAT = f(pol, n_list, d_list, th_0, lams, c_list)
        return RAT['R'], RAT['T']
        
    else:
        #case with varying thicknesses
        keys = thicks.keys()
        vals = thicks.values()
        sizes =  [len(val) for val in vals]
        params = product(*vals) #cartesian product of all thicknesses
        total = np.prod(sizes)
        # out = [np.nan]*total
        Rs = np.empty((total,n_lams))
        Ts = np.empty((total,n_lams))
        
        for i,par in enumerate(params):
            if not i%500: print(f'case {i}/{total}')
            
            for key,t in zip(keys, par):
                d_list[key] = t
            
            RAT = f(pol, n_list, d_list, th_0, lams, c_list)
            Rs[i] = RAT['R']
            Ts[i] = RAT['T']
            
        print('####finished####',)
        return Rs.reshape(*sizes,n_lams), Ts.reshape(*sizes,n_lams)


def calculate_RT_torch(stack, materials, lams, weights=None, pol='s', th_0=0.0, num_starts=50, num_epochs=150, lr=2.0, use_cuda=False):
    """
    Optimiza el espesor de un stack de capas delgadas usando PyTorch Autograd.
    
    Parámetros:
    - stack: Lista de capas del stack óptico [[d_bounds/d_fijo, 'material', 'c'/'i'], ...].
    - materials: Diccionario con funciones interpoladoras para cada material.
    - lams: Vector o Tensor con las longitudes de onda [nm].
    - weights: Tensor opcional con los pesos (por ejemplo, espectro solar * IQE * wavelength)
               para realizar una optimización por reflectancia ponderada (Jsc).
    - pol: Polarización, 's' o 'p'.
    - th_0: Ángulo de incidencia (radianes).
    - num_starts: Cantidad de semillas aleatorias para evitar mínimos locales.
    - num_epochs: Número de iteraciones del Descenso de Gradiente.
    - lr: Learning rate (tasa de aprendizaje) para el optimizador Adam.
    - use_cuda: Si es True, utiliza la GPU para acelerar los cálculos (si está disponible).
    """
    device = torch.device("cuda" if use_cuda and torch.cuda.is_available() else "cpu")
    print(f"Ejecutando en dispositivo: {device}")
    
    # 1. Asegurar que lams sea un tensor de PyTorch en el dispositivo
    if not isinstance(lams, torch.Tensor):
        lams_np = np.array(lams, dtype=np.float64)
        lams = torch.tensor(lams_np, dtype=torch.float64, device=device)
    else:
        lams_np = lams.cpu().numpy()
        lams = lams.to(device)

    # 2. Generar n_list, d_list, c_list a partir de stack y materials
    n_list_np, d_list, c_list = stack2tmm(stack, materials, lams_np, add_inf=False)
    
    # Convertir n_list (lista de arrays NumPy) a Tensor de PyTorch
    n_list = torch.tensor(np.array(n_list_np), dtype=torch.complex128, device=device)
    
    # d_bounds corresponde a las capas finitas internas
    d_bounds = d_list[1:-1]
    
    if weights is not None:
        if not isinstance(weights, torch.Tensor):
            weights = torch.tensor(weights, dtype=torch.float64, device=device)
        else:
            weights = weights.to(device)
        # Normalizar pesos para asegurar promedio ponderado correcto
        weights = weights / weights.sum()
    
    num_layers = n_list.shape[0]
    num_finite_layers = num_layers - 2
    num_wl = n_list.shape[1]
    
    if len(d_bounds) != num_finite_layers:
        raise ValueError(f"d_bounds debe tener longitud igual al número de capas finitas ({num_finite_layers}), pero tiene {len(d_bounds)}.")
    
    # 1. ANALIZAR CAPAS OPTIMIZABLES Y FIJAS
    is_optimizable = []
    opt_bounds = []  # Límites (min, max) para optimizables
    fixed_values = {}  # Mapa de idx -> valor_fijo para no optimizables
    
    for idx, b in enumerate(d_bounds):
        if isinstance(b, (tuple, list)):
            d_min, d_max = float(b[0]), float(b[1])
            if d_min == d_max:
                is_optimizable.append(False)
                fixed_values[idx] = d_min
            else:
                is_optimizable.append(True)
                opt_bounds.append((d_min, d_max))
        else:
            is_optimizable.append(False)
            fixed_values[idx] = float(b)
            
    num_opt_layers = sum(is_optimizable)
    print(f"Capas a optimizar: {num_opt_layers} | Capas fijas: {num_finite_layers - num_opt_layers}")
    
    use_inc = 'i' in c_list[1:-1] if c_list is not None else False

    if num_opt_layers == 0:
        print("Todas las capas tienen espesor fijo. Calculando reflectancia directa sin optimización...")
        d_fisico = torch.tensor([[fixed_values[i] for i in range(num_finite_layers)]], dtype=torch.float64, device=device)
        inf_col_1 = torch.full((1, 1), float('inf'), dtype=torch.float64, device=device)
        d_full = torch.cat([inf_col_1, d_fisico, inf_col_1], dim=1)
        d_full_3d = d_full.unsqueeze(-1).expand(-1, -1, num_wl)
        
        with torch.no_grad():
            if not use_inc:
                result = tmm.coh_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            else:
                result = tmm.inc_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d, c_list=c_list, th_0=th_0, lam_vac=lams)
            R_final = result['R']  # [1, num_wl]
        
        if weights is None:
            loss_val = R_final.mean(dim=-1).item()
            print(f"Reflectancia media: {loss_val*100:.2f}%")
        else:
            loss_val = (R_final * weights).sum(dim=-1).item()
            print(f"Reflectancia media ponderada: {loss_val*100:.2f}%")
            
        best_thicknesses = d_fisico[0].cpu().numpy()
        best_R_curve = R_final[0].cpu().numpy()
        return best_thicknesses, best_R_curve
        
    # 2. INICIALIZACIÓN MULTI-START EN EL ESPACIO NO ACOTODO (LOGIT)
    d_initial = torch.zeros((num_starts, num_opt_layers), dtype=torch.float64, device=device)
    opt_idx = 0
    for idx, opt in enumerate(is_optimizable):
        if opt:
            d_min, d_max = opt_bounds[opt_idx]
            d_init_phys = torch.empty(num_starts, device=device).uniform_(d_min, d_max)
            # Mapeo inverso de sigmoide (logit)
            p = (d_init_phys - d_min) / (d_max - d_min)
            p = torch.clamp(p, 1e-7, 1.0 - 1e-7)  # Evitar desbordes numéricos
            d_initial[:, opt_idx] = torch.log(p / (1.0 - p))
            opt_idx += 1
            
    d_opt = d_initial.clone().detach().requires_grad_(True)
    
    # 3. OPTIMIZADOR
    optimizer = optim.Adam([d_opt], lr=lr)
    
    print(f"Iniciando optimización con Autograd ({num_starts} semillas en paralelo)...")
    
    # Columna de infinitos para sustrato/superestrato: [num_starts, 1]
    inf_col = torch.full((num_starts, 1), float('inf'), dtype=torch.float64, device=device)
    
    # Helper para construir d_fisico a partir de d_opt y las capas fijas
    def reconstruct_d_fisico(d_opt_tensor):
        d_fisico_cols = []
        o_idx = 0
        for idx, opt in enumerate(is_optimizable):
            if opt:
                d_min, d_max = opt_bounds[o_idx]
                # Mapeo sigmoide: garantiza d_min <= espesor <= d_max strictly
                col = d_min + (d_max - d_min) * torch.sigmoid(d_opt_tensor[:, o_idx])
                d_fisico_cols.append(col)
                o_idx += 1
            else:
                val = fixed_values[idx]
                col = torch.full((d_opt_tensor.shape[0],), val, dtype=torch.float64, device=device)
                d_fisico_cols.append(col)
        return torch.stack(d_fisico_cols, dim=1)
    
    use_inc = 'i' in c_list[1:-1] if c_list is not None else False
    
    # 4. BUCLE DE OPTIMIZACIÓN (sin loop sobre semillas)
    for epoch in range(num_epochs):
        optimizer.zero_grad()
        
        d_fisico = reconstruct_d_fisico(d_opt)  # [num_starts, num_finite_layers]
        
        # Construir d_list BATCHED: [num_starts, num_layers]
        d_full = torch.cat([inf_col, d_fisico, inf_col], dim=1)
        # Expandir a [num_starts, num_layers, num_wl]
        d_full_3d = d_full.unsqueeze(-1).expand(-1, -1, num_wl)
        
        # UNA SOLA llamada para TODAS las semillas
        if not use_inc:
            result = tmm.coh_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d,
                                        th_0=th_0, lam_vac=lams)
        else:
            result = tmm.inc_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d, c_list=c_list,
                                        th_0=th_0, lam_vac=lams)
            
        R = result['R']  # [num_starts, num_wl]
        
        # Loss (promedio simple o ponderado por weights)
        if weights is None:
            loss_per_start = R.mean(dim=-1)
        else:
            loss_per_start = (R * weights).sum(dim=-1)
            
        loss = loss_per_start.sum()
        
        loss.backward()
        optimizer.step()
        
        if epoch % 50 == 0:
            mejor_R_actual = loss_per_start.min().item()
            if weights is None:
                print(f"Epoch {epoch}/{num_epochs} | Mejor Reflectancia Media: {mejor_R_actual*100:.2f}%")
            else:
                print(f"Epoch {epoch}/{num_epochs} | Mejor Reflectancia Media Ponderada: {mejor_R_actual*100:.2f}%")

    # 5. EXTRACCIÓN DEL MÍNIMO GLOBAL
    with torch.no_grad():
        d_final_fisico = reconstruct_d_fisico(d_opt)
        d_full = torch.cat([inf_col, d_final_fisico, inf_col], dim=1)
        d_full_3d = d_full.unsqueeze(-1).expand(-1, -1, num_wl)
        if not use_inc:
            result = tmm.coh_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d,
                                        th_0=th_0, lam_vac=lams)
        else:
            result = tmm.inc_tmm_torch(pol=pol, n_list=n_list, d_list=d_full_3d, c_list=c_list,
                                        th_0=th_0, lam_vac=lams)
        R_final = result['R']
    
    if weights is None:
        loss_final = R_final.mean(dim=-1)
    else:
        loss_final = (R_final * weights).sum(dim=-1)
        
    best_idx = loss_final.argmin()
    
    best_thicknesses = d_final_fisico[best_idx].cpu().detach().numpy()
    best_R_curve = R_final[best_idx].cpu().detach().numpy()
    
    print(f"\n¡Optimización completada!")
    print(f"Espesores óptimos encontrados: {best_thicknesses} nm")
    if weights is None:
        print(f"Reflectancia media óptima: {loss_final[best_idx].item()*100:.2f}%")
    else:
        print(f"Reflectancia media ponderada óptima: {loss_final[best_idx].item()*100:.2f}%")
    
    return best_thicknesses, best_R_curve


def coh_tmm_torch_batched(pol, n_list, d_list, th_0, lam_vac):
    """
    Versión diferenciable y paralelizada (doblemente BATCHED: por semillas y lambdas) 
    del Método de la Matriz de Transferencia en PyTorch.
    Soporta que n_list y d_list tengan forma [batch, num_layers, num_wl].
    """
    device = n_list.device
    batch_size, num_layers, num_wl = n_list.shape
    
    if not isinstance(lam_vac, torch.Tensor):
        lam_vac = torch.tensor(lam_vac, dtype=torch.float64, device=device)
    else:
        lam_vac = lam_vac.to(device)
        
    if not isinstance(th_0, torch.Tensor):
        th_0 = torch.tensor(th_0, dtype=n_list.real.dtype, device=device)
    else:
        th_0 = th_0.to(device)
        
    sin_th_0 = torch.sin(th_0)
    # th_list shape: [batch, num_layers, num_wl]
    # Slicing n_list[:, 0:1, :] obtiene el superestrato (capa 0) para cada lote
    th_list = torch.asin(n_list[:, 0:1, :] * sin_th_0 / n_list)
    
    # kz_list: [batch, num_layers, num_wl]
    kz_list = 2 * torch.pi * n_list * torch.cos(th_list) / lam_vac.view(1, 1, num_wl)
    
    # delta: [batch, num_layers, num_wl]
    delta = kz_list * d_list
    
    # Prevenir desbordamiento por absorción alta en capas internas finitas
    inner = delta[:, 1:-1, :]
    cond = inner.imag > 100
    inner = torch.where(cond, inner.real + 100j, inner)
    delta = torch.cat([delta[:, :1, :], inner, delta[:, -1:, :]], dim=1)
    
    t_list_vals = []
    r_list_vals = []
    
    for i in range(num_layers - 1):
        n_i = n_list[:, i, :]
        n_f = n_list[:, i+1, :]
        th_i = th_list[:, i, :]
        th_f = th_list[:, i+1, :]
        
        cos_th_i = torch.cos(th_i)
        cos_th_f = torch.cos(th_f)
        
        if pol == 's':
            r_val = (n_i * cos_th_i - n_f * cos_th_f) / (n_i * cos_th_i + n_f * cos_th_f)
            t_val = 2 * n_i * cos_th_i / (n_i * cos_th_i + n_f * cos_th_f)
        elif pol == 'p':
            r_val = (n_f * cos_th_i - n_i * cos_th_f) / (n_f * cos_th_i + n_i * cos_th_f)
            t_val = 2 * n_i * cos_th_i / (n_f * cos_th_i + n_i * cos_th_f)
        else:
            raise ValueError("Polarization must be 's' or 'p'")
            
        r_list_vals.append(r_val)
        t_list_vals.append(t_val)
        
    # Inicializar la matriz de transferencia Mtilde = [batch, num_wl, 2, 2]
    ones_bw = torch.ones(batch_size, num_wl, dtype=delta.dtype, device=device)
    zeros_bw = torch.zeros_like(ones_bw)
    
    Mtilde = torch.cat([
        ones_bw.unsqueeze(-1), zeros_bw.unsqueeze(-1),
        zeros_bw.unsqueeze(-1), ones_bw.unsqueeze(-1)
    ], dim=-1).view(batch_size, num_wl, 2, 2)
    
    for i in range(1, num_layers - 1):
        delta_i = delta[:, i, :]
        r_i = r_list_vals[i]
        t_i = t_list_vals[i]
        
        exp_minus = torch.exp(-1j * delta_i)
        exp_plus = torch.exp(1j * delta_i)
        zeros_i = torch.zeros_like(delta_i)
        ones_i = torch.ones_like(delta_i)
        
        # A y B shape: [batch, num_wl, 2, 2]
        A = torch.cat([
            exp_minus.unsqueeze(-1), zeros_i.unsqueeze(-1),
            zeros_i.unsqueeze(-1), exp_plus.unsqueeze(-1)
        ], dim=-1).view(batch_size, num_wl, 2, 2)
        
        B = torch.cat([
            ones_i.unsqueeze(-1), r_i.unsqueeze(-1),
            r_i.unsqueeze(-1), ones_i.unsqueeze(-1)
        ], dim=-1).view(batch_size, num_wl, 2, 2)
        
        d_coeff = (1.0 / t_i).unsqueeze(-1).unsqueeze(-1)
        M_i = d_coeff * torch.matmul(A, B)
        Mtilde = torch.matmul(Mtilde, M_i)
        
    # Interfaz frontal (0 -> 1)
    r_0 = r_list_vals[0]
    t_0 = t_list_vals[0]
    ones_0 = torch.ones_like(r_0)
    
    A_front = torch.cat([
        ones_0.unsqueeze(-1), r_0.unsqueeze(-1),
        r_0.unsqueeze(-1), ones_0.unsqueeze(-1)
    ], dim=-1).view(batch_size, num_wl, 2, 2)
    
    d_front = (1.0 / t_0).unsqueeze(-1).unsqueeze(-1)
    M_front = d_front * A_front
    Mtilde = torch.matmul(M_front, Mtilde)
    
    r = Mtilde[:, :, 1, 0] / Mtilde[:, :, 0, 0]
    t = 1.0 / Mtilde[:, :, 0, 0]
    
    return {'r': r, 't': t}




def fit_ellipsometry_torch(n_list, d_bounds, lams, Is_exp, Ic_exp, Data_R=None, th_0=0.0, 
                           num_starts=50, num_epochs=150, lr=2.0, use_cuda=False, 
                           layer_models=None, layer_names=None, std_Is=None, std_Ic=None,
                           check_cancel_fn=None, n_max_limits=None):
    """
    Optimiza el espesor de un stack de capas delgadas y opcionalmente los parámetros 
    de modelos de dispersión (como Cauchy) usando PyTorch Autograd.
    
    Parámetros:
    - n_list: Tensor complejo de PyTorch con los índices [num_capas, num_lams] o lista de índices/tensores.
    - d_bounds: Lista de límites para los espesores de capas finitas.
    - lams: Tensor con las longitudes de onda.
    - Is_exp: Tensor PyTorch con Is experimental.
    - Ic_exp: Tensor PyTorch con Ic experimental.
    - Data_R: Ruta opcional a un archivo de texto con mediciones de reflectancia (a 0 grados).
    - th_0: Ángulo de incidencia en radianes.
    - num_starts: Cantidad de semillas aleatorias paralelas.
    - num_epochs: Número de épocas del optimizador.
    - lr: Tasa de aprendizaje.
    - use_cuda: Usar GPU si está disponible.
    - layer_models: Lista de modelos de dispersión por capa. e.g. [None, 'cauchy', None].
      Soporta 'cauchy', 'cauchy_absorbent' o un diccionario con 'model', 'bounds', 'initial'.
      Para Bruggeman: 'bruggeman' (f_air=0.5 por defecto),
        {'model': 'bruggeman', 'f_air': 0.42} (fracción fija),
        {'model': 'bruggeman', 'f_bounds': (0.3, 0.7)} (fracción optimizable).
    - layer_names: Lista opcional con los nombres de las capas.
    - n_max_limits: Tupla/lista opcional con límites máximos para el índice n en cada capa (e.g. (None, 2.5, None)).
      Si para una semilla alguna λ produce n(λ) > límite en una capa con modelo de Cauchy, esa semilla se toma como inválida.
    """
    device = torch.device("cuda" if use_cuda and torch.cuda.is_available() else "cpu")
    if use_cuda and not torch.cuda.is_available():
        print("Ejecutando en dispositivo: cpu (CUDA no disponible en el sistema)")
    else:
        print(f"Ejecutando en dispositivo: {device}")
    
    if std_Is is not None:
        if not isinstance(std_Is, torch.Tensor):
            std_Is = torch.tensor(std_Is, dtype=torch.float64, device=device)
        else:
            std_Is = std_Is.to(device)
    if std_Ic is not None:
        if not isinstance(std_Ic, torch.Tensor):
            std_Ic = torch.tensor(std_Ic, dtype=torch.float64, device=device)
        else:
            std_Ic = std_Ic.to(device)
    
    # Asegurar que lams, Is_exp, Ic_exp estén en el dispositivo
    if not isinstance(lams, torch.Tensor):
        lams = torch.tensor(lams, dtype=torch.float64, device=device)
    else:
        lams = lams.to(device)
    
    if Data_R is not None:
        R_exp_raw = np.loadtxt(Data_R)
        wl_R = np.linspace(190, 900, len(R_exp_raw))
        # bounds_error=False y fill_value="extrapolate" para evitar caídas por fuera de rango
        R_interp = interp1d(wl_R, R_exp_raw, kind='linear', bounds_error=False, fill_value="extrapolate")
        
        # Mover a CPU si lams es tensor de PyTorch (interp1d requiere numpy)
        lams_np = lams.cpu().numpy() if isinstance(lams, torch.Tensor) else np.array(lams)
        R_exp = R_interp(lams_np)
        
        # Si la reflectancia está en porcentaje (0-100), la dividimos por 100 para que coincida con el rango [0,1]
        if np.max(R_exp) > 1.0:
            print("Advertencia: Se detectó reflectancia > 1.0 en los datos experimentales. Se asume escala 0-100% y se divide por 100.")
            R_exp = R_exp / 100.0
            
        R_exp = torch.tensor(R_exp, dtype=torch.float64, device=device)
        
    if not isinstance(Is_exp, torch.Tensor):
        Is_exp = torch.tensor(Is_exp, dtype=torch.float64, device=device)
    else:
        Is_exp = Is_exp.to(device)
        
    if not isinstance(Ic_exp, torch.Tensor):
        Ic_exp = torch.tensor(Ic_exp, dtype=torch.float64, device=device)
    else:
        Ic_exp = Ic_exp.to(device)
        
    # Convertir n_list a tensor en dispositivo si es un tensor, o lista de tensores
    if isinstance(n_list, torch.Tensor):
        n_list = n_list.to(device)
        num_layers = n_list.shape[0]
        num_wl = n_list.shape[1]
    else:
        # Si es una lista, convertir tensores individuales
        n_list_tensors = []
        for item in n_list:
            if isinstance(item, torch.Tensor):
                n_list_tensors.append(item.to(device))
            else:
                # Es un dummy o placeholder para material paramétrico
                n_list_tensors.append(torch.ones_like(lams, dtype=torch.complex128, device=device))
        n_list = torch.stack(n_list_tensors, dim=0)
        num_layers = n_list.shape[0]
        num_wl = n_list.shape[1]
        
    num_finite_layers = num_layers - 2
    
    if len(d_bounds) != num_finite_layers:
        raise ValueError(f"d_bounds debe tener longitud {num_finite_layers}, pero tiene {len(d_bounds)}.")
        
    # 1. ANALIZAR MODELOS DE DISPERSIÓN PARAMÉTRICOS
    is_parametric = False
    is_parametric_optimizable = False
    disp_layers = []
    total_disp_params = 0      # Todos los parámetros (fijos + optimizables)
    total_opt_disp_params = 0  # Solo los parámetros optimizables (para p_opt)
    
    if layer_models is not None:
        for idx, model in enumerate(layer_models):
            if model is not None:
                is_parametric = True
                
                if isinstance(model, dict):
                    model_type = model.get('model', 'cauchy')
                    bounds = model.get('bounds', None)
                    initial = model.get('initial', None)
                else:
                    model_type = model
                    bounds = None
                    initial = None
                
                # Definir cantidad de parámetros y valores por defecto
                if model_type == 'cauchy':
                    num_p = 3
                    default_bounds = [(0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0)]
                    default_initial = [3.0, 0.0, 0.0]
                elif model_type == 'cauchy_absorbent':
                    num_p = 6
                    default_bounds = [(0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0), (0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0)]
                    default_initial = [3.0, 0.0, 0.0, 0.0, 0.0, 0.0]
                elif model_type == 'bruggeman':
                    if isinstance(model, dict):
                        f_air_fixed = model.get('f_air', None)
                        f_bounds = model.get('f_bounds', None)
                    else:
                        f_air_fixed = None
                        f_bounds = None
                    
                    if f_bounds is not None:
                        num_p = 1
                        if isinstance(f_bounds[0], (list, tuple)):
                            fb = f_bounds[0]
                        else:
                            fb = f_bounds
                        default_bounds = [tuple(fb)]
                        default_initial = [(fb[0] + fb[1]) / 2.0]
                    else:
                        num_p = 0
                        default_bounds = []
                        default_initial = []
                else:
                    raise ValueError(f"Modelo de dispersión no soportado: {model_type}")
                    
                if bounds is None:
                    bounds = default_bounds
                if initial is None:
                    initial = default_initial
                
                # Analizar cuáles de estos parámetros son fijos y cuáles optimizables
                params_optimizable = []
                opt_param_bounds = []
                fixed_param_values = {}
                
                for k in range(num_p):
                    p_b = bounds[k]
                    if isinstance(p_b, (list, tuple)):
                        p_min, p_max = float(p_b[0]), float(p_b[1])
                        if p_min == p_max:
                            params_optimizable.append(False)
                            fixed_param_values[k] = p_min
                        else:
                            params_optimizable.append(True)
                            opt_param_bounds.append((p_min, p_max))
                    else:
                        params_optimizable.append(False)
                        fixed_param_values[k] = float(p_b)
                
                num_opt_p = sum(params_optimizable)
                if num_opt_p > 0:
                    is_parametric_optimizable = True
                
                layer_info = {
                    'layer_idx': idx,
                    'model_type': model_type,
                    'param_bounds': bounds,
                    'param_initial': initial,
                    'num_params': num_p,
                    'params_optimizable': params_optimizable,
                    'opt_param_bounds': opt_param_bounds,
                    'fixed_param_values': fixed_param_values,
                    'num_opt_params': num_opt_p,
                    'start_idx': total_opt_disp_params  # Índice en el tensor p_opt optimizable
                }
                
                if model_type == 'bruggeman':
                    if isinstance(model, dict) and 'f_bounds' in model:
                        layer_info['f_optimizable'] = True
                    elif isinstance(model, dict) and 'f_air' in model:
                        layer_info['f_air'] = model['f_air']
                        layer_info['f_optimizable'] = False
                    else:
                        layer_info['f_air'] = 0.5
                        layer_info['f_optimizable'] = False
                
                disp_layers.append(layer_info)
                total_disp_params += num_p
                total_opt_disp_params += num_opt_p
                
    if is_parametric:
        print(f"Modelos paramétricos activos: {[d['model_type'] for d in disp_layers]} | Parámetros de dispersión: {total_opt_disp_params} optimizables, {total_disp_params - total_opt_disp_params} fijos")
    
    # 2. ANALIZAR ESPESORES OPTIMIZABLES Y FIJOS
    is_optimizable = []
    opt_bounds = []
    fixed_values = {}
    
    for idx, b in enumerate(d_bounds):
        if isinstance(b, (tuple, list)):
            d_min, d_max = float(b[0]), float(b[1])
            if d_min == d_max:
                is_optimizable.append(False)
                fixed_values[idx] = d_min #Si los espesores son iguales no se optimiza esta capa
            else:
                is_optimizable.append(True)
                opt_bounds.append((d_min, d_max))
        else:
            is_optimizable.append(False)
            fixed_values[idx] = float(b) #Si se ingresa un solo espesor no se optimiza esta capa
            
    num_opt_layers = sum(is_optimizable)
    print(f"Capas finitas a optimizar espesor: {num_opt_layers} | Capas de espesor fijo: {num_finite_layers - num_opt_layers}")
    
    # 3. INICIALIZACIÓN MULTI-START DE ESPESORES EN LOGITS
    d_initial = torch.zeros((num_starts, num_opt_layers), dtype=torch.float64, device=device)
    o_idx = 0
    for idx, opt in enumerate(is_optimizable):
        if opt:
            d_min, d_max = opt_bounds[o_idx]
            d_init_phys = torch.empty(num_starts, device=device).uniform_(d_min, d_max)
            p = (d_init_phys - d_min) / (d_max - d_min)
            p = torch.clamp(p, 1e-7, 1.0 - 1e-7)
            d_initial[:, o_idx] = torch.log(p / (1.0 - p))
            o_idx += 1
            
    d_opt = d_initial.clone().detach().requires_grad_(True)
    
    # 4. INICIALIZACIÓN MULTI-START DE PARÁMETROS DE DISPERSIÓN EN LOGITS
    if is_parametric_optimizable:
        p_initial = torch.zeros((num_starts, total_opt_disp_params), dtype=torch.float64, device=device)
        for d_lay in disp_layers:
            start_idx = d_lay['start_idx']
            o_p_idx = 0
            for k in range(d_lay['num_params']):
                if d_lay['params_optimizable'][k]:
                    p_min, p_max = d_lay['opt_param_bounds'][o_p_idx]
                    p_init_phys = torch.empty(num_starts, device=device).uniform_(p_min, p_max)
                    p = (p_init_phys - p_min) / (p_max - p_min)
                    p = torch.clamp(p, 1e-7, 1.0 - 1e-7)
                    p_initial[:, start_idx + o_p_idx] = torch.log(p / (1.0 - p))
                    o_p_idx += 1
                
        p_opt = p_initial.clone().detach().requires_grad_(True)
        
    # Filtrar tensores que tengan parámetros reales para optimizar
    params_to_opt = []
    if num_opt_layers > 0:
        params_to_opt.append(d_opt)
    if is_parametric_optimizable:
        params_to_opt.append(p_opt)
        
    if len(params_to_opt) > 0:
        optimizer = optim.Adam(params_to_opt, lr=lr)
    else:
        # Mock/Dummy Optimizer si no hay variables optimizables
        class DummyOptimizer:
            def zero_grad(self): pass
            def step(self): pass
        optimizer = DummyOptimizer()
        
    print(f"Iniciando optimización elipsométrica con Autograd ({num_starts} semillas en paralelo)...")
    
    inf_col = torch.full((num_starts, 1), float('inf'), dtype=torch.float64, device=device)
    
    def reconstruct_d_fisico(d_opt_tensor):
        d_fisico_cols = []
        o_idx = 0
        for idx, opt in enumerate(is_optimizable):
            if opt:
                d_min, d_max = opt_bounds[o_idx]
                col = d_min + (d_max - d_min) * torch.sigmoid(d_opt_tensor[:, o_idx])
                d_fisico_cols.append(col)
                o_idx += 1
            else:
                val = fixed_values[idx]
                col = torch.full((d_opt_tensor.shape[0],), val, dtype=torch.float64, device=device)
                d_fisico_cols.append(col)
        return torch.stack(d_fisico_cols, dim=1)
    # Precomputar dependencias de lams para acelerar los cálculos paramétricos
    lams_2d = lams.unsqueeze(0)  # [1, num_wl]
    inv_lam2 = 1e4 / (lams_2d ** 2)
    inv_lam4 = 1e9 / (lams_2d ** 4)

    def reconstruct_n_list_batched(p_opt_tensor):
        # Construir una lista de tensores por capa (sin operaciones in-place)
        layer_tensors = [None] * num_layers
        
        # 1. Copiar capas estáticas (no paramétricas)
        for idx in range(num_layers):
            is_this_parametric = False
            for d_lay in disp_layers:
                if d_lay['layer_idx'] == idx:
                    is_this_parametric = True
                    break
            if not is_this_parametric:
                layer_tensors[idx] = n_list[idx].unsqueeze(0).expand(num_starts, -1)
                
        # 2. Calcular capas con modelos de dispersión (Cauchy, etc.)
        for d_lay in disp_layers:
            if d_lay['model_type'] in ['cauchy', 'cauchy_absorbent']:
                idx = d_lay['layer_idx']
                start_idx = d_lay['start_idx']
                model_type = d_lay['model_type']
                
                p_phys_list = []
                o_p_idx = 0
                for k in range(d_lay['num_params']):
                    if d_lay['params_optimizable'][k]:
                        p_min, p_max = d_lay['opt_param_bounds'][o_p_idx]
                        p_logit = p_opt_tensor[:, start_idx + o_p_idx]
                        p_phys = p_min + (p_max - p_min) * torch.sigmoid(p_logit)
                        o_p_idx += 1
                    else:
                        val = d_lay['fixed_param_values'][k]
                        p_phys = torch.full((num_starts,), val, dtype=torch.float64, device=device)
                    p_phys_list.append(p_phys)
                    
                if model_type == 'cauchy':
                    A = p_phys_list[0].unsqueeze(1)
                    B = p_phys_list[1].unsqueeze(1)
                    C = p_phys_list[2].unsqueeze(1)
                    n_calc = A + B * inv_lam2 + C * inv_lam4
                    n_calc = torch.clamp(n_calc, min=1.0)
                    n_complex = n_calc.to(torch.complex128)
                elif model_type == 'cauchy_absorbent':
                    A = p_phys_list[0].unsqueeze(1)
                    B = p_phys_list[1].unsqueeze(1)
                    C = p_phys_list[2].unsqueeze(1)
                    D = p_phys_list[3].unsqueeze(1)
                    E = p_phys_list[4].unsqueeze(1)
                    F = p_phys_list[5].unsqueeze(1)
                    n_calc = A + B * inv_lam2 + C * inv_lam4
                    n_calc = torch.clamp(n_calc, min=1.0)
                    k_calc = D + E * inv_lam2 + F * inv_lam4
                    k_calc = torch.clamp(k_calc, min=0.0)
                    n_complex = torch.complex(n_calc, k_calc)
                    
                layer_tensors[idx] = n_complex

        # 3. Capas dependientes (Bruggeman EMA) - usan el resultado de Cauchy
        for d_lay in disp_layers:
            if d_lay['model_type'] == 'bruggeman':
                idx = d_lay['layer_idx']
                # Tomamos el tensor de la capa adyacente (idx + 1), ya calculada
                n_base = layer_tensors[idx + 1]
                # Tensor de aire (n=1.0) con la misma forma que n_base
                n_air = torch.ones_like(n_base)
                
                # Determinar f_air: fijo o desde parámetros optimizables
                if d_lay.get('f_optimizable', False) and d_lay['params_optimizable'][0]:
                    start_idx = d_lay['start_idx']
                    p_min, p_max = d_lay['opt_param_bounds'][0]
                    p_logit = p_opt_tensor[:, start_idx]
                    f_air = p_min + (p_max - p_min) * torch.sigmoid(p_logit)
                    f_air = f_air.unsqueeze(1)  # [num_starts, 1] para broadcasting
                else:
                    if d_lay.get('f_optimizable', False):
                        f_air = d_lay['fixed_param_values'][0]
                    else:
                        f_air = d_lay.get('f_air', 0.5)
                
                # Bruggeman directamente sobre tensores (mantiene Autograd)
                layer_tensors[idx] = n_eff_torch(n_base, n_air, f_air)

        # Apilar en tensor 3D sin operaciones in-place
        return torch.stack(layer_tensors, dim=1)

    # Registrar el mínimo histórico de cada semilla (semilla -> mejor pérdida, logits)
    best_losses_per_start = torch.full((num_starts,), float('inf'), dtype=torch.float64, device=device)
    best_d_opt_data = d_opt.clone().detach()
    if is_parametric_optimizable:
        best_p_opt_data = p_opt.clone().detach()

    # 5. BUCLE DE OPTIMIZACIÓN
    for epoch in range(num_epochs):
        if check_cancel_fn is not None:
            check_cancel_fn()
            
        optimizer.zero_grad()
        
        d_fisico = reconstruct_d_fisico(d_opt)
        d_full = torch.cat([inf_col, d_fisico, inf_col], dim=1)
        d_full_3d = d_full.unsqueeze(-1).expand(-1, -1, num_wl)
        
        if is_parametric:
            n_list_3d = reconstruct_n_list_batched(p_opt if is_parametric_optimizable else None)
            res_s = coh_tmm_torch_batched(pol='s', n_list=n_list_3d, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            res_p = coh_tmm_torch_batched(pol='p', n_list=n_list_3d, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            if Data_R is not None:
                res_0deg = coh_tmm_torch_batched(pol='s', n_list=n_list_3d, d_list=d_full_3d, th_0=0.0, lam_vac=lams)
        else:
            res_s = tmm.coh_tmm_torch(pol='s', n_list=n_list, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            res_p = tmm.coh_tmm_torch(pol='p', n_list=n_list, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            if Data_R is not None:
                res_0deg = tmm.coh_tmm_torch(pol='s', n_list=n_list, d_list=d_full_3d, th_0=0.0, lam_vac=lams)
            
        r_s = res_s['r']
        r_p = res_p['r']
        
        if Data_R is not None:
            R_teo = (res_0deg['r'] * res_0deg['r'].conj()).real
            error_R = (R_teo - R_exp) ** 2
        
        rho = torch.conj(r_p / r_s)
        psi_teo = torch.atan(torch.abs(rho))
        delta_teo = torch.angle(rho)
        
        Is_teo = torch.sin(2 * psi_teo) * torch.sin(delta_teo)
        Ic_teo = torch.sin(2 * psi_teo) * torch.cos(delta_teo)
        
        # Calcular grados de libertad (dof) para Chi2 Reducido
        num_free_params = num_opt_layers
        if is_parametric:
            num_free_params += total_opt_disp_params
        dof = 2 * num_wl - num_free_params
        if dof < 1:
            dof = 1
            
        if std_Is is not None and std_Ic is not None:
            error_Is = ((Is_teo - Is_exp) / std_Is) ** 2
            error_Ic = ((Ic_teo - Ic_exp) / std_Ic) ** 2
            if Data_R is not None:
                loss_per_start = (error_Is + error_Ic).sum(dim=-1) / dof + error_R.mean(dim=-1)
            else:
                loss_per_start = (error_Is + error_Ic).sum(dim=-1) / dof
        else:
            error_Is = (Is_teo - Is_exp) ** 2
            error_Ic = (Ic_teo - Ic_exp) ** 2
            if Data_R is not None:
                loss_per_start = (error_Is + error_Ic + error_R).mean(dim=-1)
            else:
                loss_per_start = (error_Is + error_Ic).mean(dim=-1)
            
        loss = loss_per_start.sum()
        
        if len(params_to_opt) > 0:
            loss.backward()
            optimizer.step()
        
        # Prevenir saturación de sigmoid: clampear logits
        with torch.no_grad():
            d_opt.data.clamp_(-5, 5)
            if is_parametric_optimizable:
                p_opt.data.clamp_(-5, 5)
                
            # Actualizar el histórico del mínimo real por semilla
            improved = loss_per_start < best_losses_per_start
            if improved.any():
                best_losses_per_start[improved] = loss_per_start[improved]
                best_d_opt_data[improved] = d_opt.data[improved]
                if is_parametric_optimizable:
                    best_p_opt_data[improved] = p_opt.data[improved]
        
        if epoch % 50 == 0 or epoch == num_epochs - 1:
            mejor_loss = loss_per_start.min().item()
            mejor_historico = best_losses_per_start.min().item()
            metric_name = "Chi2 Red" if (std_Is is not None and std_Ic is not None) else "MSE"
            print(f"Epoch {epoch}/{num_epochs} | Mejor Loss Actual ({metric_name}): {mejor_loss:.6f} | Mínimo Histórico: {mejor_historico:.6f}")
            
    # 6. EXTRACCIÓN DEL MÍNIMO GLOBAL (USANDO LOS LOGITS HISTÓRICOS MÍNIMOS)
    with torch.no_grad():
        d_final_fisico = reconstruct_d_fisico(best_d_opt_data)
        d_full = torch.cat([inf_col, d_final_fisico, inf_col], dim=1)
        d_full_3d = d_full.unsqueeze(-1).expand(-1, -1, num_wl)
        
        if is_parametric:
            n_list_3d_final = reconstruct_n_list_batched(best_p_opt_data if is_parametric_optimizable else None)
            res_s_final = coh_tmm_torch_batched(pol='s', n_list=n_list_3d_final, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            res_p_final = coh_tmm_torch_batched(pol='p', n_list=n_list_3d_final, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
        else:
            res_s_final = tmm.coh_tmm_torch(pol='s', n_list=n_list, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            res_p_final = tmm.coh_tmm_torch(pol='p', n_list=n_list, d_list=d_full_3d, th_0=th_0, lam_vac=lams)
            
        r_s_final = res_s_final['r']
        r_p_final = res_p_final['r']
        
        rho_final = torch.conj(r_p_final / r_s_final)
        psi_teo_final = torch.atan(torch.abs(rho_final))
        delta_teo_final = torch.angle(rho_final)
        
        Is_teo_final = torch.sin(2 * psi_teo_final) * torch.sin(delta_teo_final)
        Ic_teo_final = torch.sin(2 * psi_teo_final) * torch.cos(delta_teo_final)
        
    # Filtrar semillas válidas respetando n_max_limits y descartando NaNs/Infs
    valid_mask = torch.ones(num_starts, dtype=torch.bool, device=device)
    valid_mask = valid_mask & ~torch.isnan(best_losses_per_start) & ~torch.isinf(best_losses_per_start)
    
    if n_max_limits is not None and is_parametric:
        for idx in range(num_layers):
            limit = None
            if len(n_max_limits) == num_layers and idx < len(n_max_limits):
                limit = n_max_limits[idx]
            elif len(n_max_limits) == num_finite_layers and 1 <= idx <= num_finite_layers:
                limit = n_max_limits[idx - 1]
            elif idx < len(n_max_limits):
                limit = n_max_limits[idx]
                
            if limit is not None:
                limit_val = float(limit)
                n_layer_real = n_list_3d_final[:, idx, :].real
                layer_valid = (n_layer_real <= limit_val).all(dim=-1)
                valid_mask = valid_mask & layer_valid

    valid_indices = torch.where(valid_mask)[0]
    if len(valid_indices) > 0:
        valid_losses = best_losses_per_start[valid_mask]
        best_valid_sub_idx = valid_losses.argmin()
        best_idx = valid_indices[best_valid_sub_idx].item()
        if n_max_limits is not None:
            print(f"Filtrado n_max: {len(valid_indices)}/{num_starts} semillas cumplen las condiciones de n_max <= {n_max_limits}.")
    else:
        if n_max_limits is not None:
            print(f"Advertencia: Ninguna de las {num_starts} semillas cumplió con el límite n_max={n_max_limits}. Se seleccionará la mejor semilla sin aplicar el filtro.")
        best_idx = best_losses_per_start.argmin().item()
    
    best_thicknesses = d_final_fisico[best_idx].cpu().detach().numpy()
    best_Is_curve = Is_teo_final[best_idx].cpu().detach().numpy()
    best_Ic_curve = Ic_teo_final[best_idx].cpu().detach().numpy()
    
    best_params = {}
    best_params['chi2_min'] = float(best_losses_per_start[best_idx].item())
    
    if is_parametric:
        with torch.no_grad():
            for d_lay in disp_layers:
                idx = d_lay['layer_idx']
                start_idx = d_lay['start_idx']
                bounds = d_lay['param_bounds']
                model_type = d_lay['model_type']
                
                if layer_names is not None and idx < len(layer_names):
                    layer_name = layer_names[idx]
                else:
                    layer_name = f"layer_{idx}"
                    
                p_best_phys = []
                o_p_idx = 0
                for k in range(d_lay['num_params']):
                    if d_lay['params_optimizable'][k]:
                        p_logit = best_p_opt_data[best_idx, start_idx + o_p_idx]
                        p_min, p_max = d_lay['opt_param_bounds'][o_p_idx]
                        p_phys = p_min + (p_max - p_min) * torch.sigmoid(p_logit)
                        p_best_phys.append(p_phys.item())
                        o_p_idx += 1
                    else:
                        val = d_lay['fixed_param_values'][k]
                        p_best_phys.append(val)
                        
                if model_type == 'cauchy':
                    layer_params = {
                        'model': 'cauchy',
                        'A': p_best_phys[0],
                        'B': p_best_phys[1],
                        'C': p_best_phys[2]
                    }
                elif model_type == 'cauchy_absorbent':
                    layer_params = {
                        'model': 'cauchy_absorbent',
                        'A': p_best_phys[0],
                        'B': p_best_phys[1],
                        'C': p_best_phys[2],
                        'D': p_best_phys[3],
                        'E': p_best_phys[4],
                        'F': p_best_phys[5]
                    }
                elif model_type == 'bruggeman':
                    brugg_info = {
                        'model': 'bruggeman_EMA',
                        'linked_to': layer_names[idx+1] if layer_names else f'layer_{idx+1}',
                    }
                    if d_lay.get('f_optimizable', False) and d_lay['params_optimizable'][0]:
                        p_logit_f = best_p_opt_data[best_idx, start_idx]
                        p_min_f, p_max_f = d_lay['opt_param_bounds'][0]
                        f_opt = p_min_f + (p_max_f - p_min_f) * torch.sigmoid(p_logit_f)
                        brugg_info['f_air'] = f_opt.item()
                        brugg_info['f_optimized'] = True
                    else:
                        if d_lay.get('f_optimizable', False):
                            brugg_info['f_air'] = d_lay['fixed_param_values'][0]
                        else:
                            brugg_info['f_air'] = d_lay.get('f_air', 0.5)
                        brugg_info['f_optimized'] = False
                    layer_params = brugg_info

                best_params[layer_name] = layer_params
                best_params[idx] = layer_params
                    
    print(f"\n¡Optimización elipsométrica completada!")
    print(f"Espesores óptimos encontrados: {best_thicknesses} nm")
    metric_name = "Chi2 Red" if (std_Is is not None and std_Ic is not None) else "MSE"
    print(f"Error mínimo ({metric_name}): {best_losses_per_start[best_idx].item():.6f}")
    if is_parametric:
        print("Parámetros de dispersión óptimos:")
        for name, params in best_params.items():
            if name == 'chi2_min' or isinstance(name, int):
                continue
            print(f"  Material: {name} ({params['model']})")
            for k, val in params.items():
                if k != 'model':
                    if isinstance(val, (int, float)):
                        print(f"    {k} = {val:.6f}")
                    else:
                        print(f"    {k} = {val}")
                    
    if layer_models is not None:
        return best_thicknesses, best_Is_curve, best_Ic_curve, best_params
    else:
        return best_thicknesses, best_Is_curve, best_Ic_curve

def load_stack(filename, materials):
    '''Loads stack from Xlsx used by IR-S in Matlab. 'Materials' is a dict 
    where new materials will be added. Only materials not already present will 
    be added. YOU SHOULD CHECK that different nk's have different material names
    '''
    
    wb = load_workbook(filename = filename)
    ws = wb['MJSC Definition']
    
    stack = []
    start_row = 19
    end_row = 35
    
    for row in range(start_row, end_row+1):
        
        mat = ws.cell(row=row, column=3).value #C
        nk = ws.cell(row=row, column=5).value  #E
        thick = ws.cell(row=row, column=7).value #G
        
        if mat not in materials:
            # print('working on '+mat)
            materials[mat] = (load_interp('./nkdata/optical/' + nk))
        
        stack.append([thick, mat, 'c'])

    return  stack, materials

def plot_js(titulo, e_porosa, e_densa, js,espesores_comparación = None,label_comparación = "experimental",label_simulado = "optimo simulado"):
    print(f'caso: {titulo}')
    X, Y = np.meshgrid(e_porosa, e_densa)
    maxj = js.max()
    maxind = np.unravel_index(js.argmax(),js.shape)
    x_opt = X[maxind]
    y_opt = Y[maxind]
    print(f'e_superior: {x_opt}, e_inferior:{y_opt}')
    print(f'Max Jsc = {maxj:.4f}\n')
    Z = 100*js/maxj
    

    fig,ax=plt.subplots(1,1)
    colors =['blue','navy','indigo','purple','red','orangered','orange','w']
    cp = ax.contourf(X, Y, Z,levels=[70,80,85,90,95,98,99, 100], colors=colors)
    cbar =fig.colorbar(cp) # Add a colorbar to a plot
    # 8/4 le agregue este label
    cbar.set_label('Jsc_Normalizada [%]', rotation=270, labelpad=15)
    # 8/4 le agregue este scatter
    ax.scatter(x_opt, y_opt, color='white', edgecolors='black', s=100, marker='*', label=f'{label_simulado}')
    #ahora el valor experimental de Simon
    if espesores_comparación != None: 
        espesores_Simon = espesores_comparación
        ax.scatter(*espesores_Simon, color='green', edgecolors='black', s=100, marker='*', label=f'{label_comparación}')
    
    ax.set_title(titulo)
    ax.set_xlabel('Espesor superior [nm]')
    ax.set_ylabel('Espesor inferior [nm]')
    ax.legend()
    plt.show()

def autorange_fit_ellipsometry_torch(n_list, d_bounds, lams, Is_exp, Ic_exp, Data_R=None, th_0=0.0, 
                                     num_starts=50, num_epochs=150, lr=2.0, use_cuda=False, 
                                     layer_models=None, layer_names=None, std_Is=None, std_Ic=None,
                                     n_max_limits=None,
                                     tolerance=0.01, max_attempts=10, expansion_factor=0.25, contraction_factor=0.1,
                                     check_cancel_fn=None):
    """
    Optimiza el espesor y los parámetros de dispersión de forma recursiva (Autorange).
    Si detecta saturación en los límites de rango de algún parámetro, amplía el rango en esa dirección
    y reduce el opuesto, volviendo a ejecutar la optimización hasta que ningún parámetro sature
    o se alcance el límite de intentos (max_attempts).
    """
    import copy
    
    # 1. Copiar y estructurar d_bounds y layer_models para poder mutarlos
    d_bounds_work = [list(b) if isinstance(b, (tuple, list)) else b for b in d_bounds]
    
    layer_models_work = []
    if layer_models is not None:
        for model in layer_models:
            if model is None:
                layer_models_work.append(None)
            else:
                if isinstance(model, str):
                    model_dict = {'model': model}
                else:
                    model_dict = copy.deepcopy(model)
                
                model_type = model_dict.get('model', 'cauchy')
                if model_type == 'cauchy':
                    if 'bounds' not in model_dict or model_dict['bounds'] is None:
                        model_dict['bounds'] = [(0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0)]
                    if 'initial' not in model_dict or model_dict['initial'] is None:
                        model_dict['initial'] = [3.0, 0.0, 0.0]
                elif model_type == 'cauchy_absorbent':
                    if 'bounds' not in model_dict or model_dict['bounds'] is None:
                        model_dict['bounds'] = [(0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0), (0.0, 4.0), (-4.0, 4.0), (-4.0, 4.0)]
                    if 'initial' not in model_dict or model_dict['initial'] is None:
                        model_dict['initial'] = [3.0, 0.0, 0.0, 0.0, 0.0, 0.0]
                elif model_type == 'bruggeman':
                    if 'f_bounds' in model_dict and model_dict['f_bounds'] is not None:
                        pass
                layer_models_work.append(model_dict)
    else:
        layer_models_work = None

    # Función auxiliar de ajuste de rangos
    def adjust_val_range(val_min, val_max, val_opt, min_abs=None, max_abs=None):
        width = val_max - val_min
        if width <= 0:
            return val_min, val_max, False
        
        # Verificar límite superior
        if (val_max - val_opt) < tolerance * width:
            val_max_new = val_max + expansion_factor * width
            if max_abs is not None:
                val_max_new = min(val_max_new, max_abs)
            # Solo contraer val_min si val_max realmente se expandió
            if val_max_new > val_max:
                val_min_new = val_min + contraction_factor * width
                if min_abs is not None:
                    val_min_new = max(val_min_new, min_abs)
            else:
                val_min_new = val_min
                
            if (val_min_new != val_min or val_max_new != val_max) and val_min_new < val_max_new:
                return val_min_new, val_max_new, True
                
        # Verificar límite inferior
        elif (val_opt - val_min) < tolerance * width:
            val_min_new = val_min - expansion_factor * width
            if min_abs is not None:
                val_min_new = max(val_min_new, min_abs)
            # Solo contraer val_max si val_min realmente se expandió
            if val_min_new < val_min:
                val_max_new = val_max - contraction_factor * width
                if max_abs is not None:
                    val_max_new = min(val_max_new, max_abs)
            else:
                val_max_new = val_max
                
            if (val_min_new != val_min or val_max_new != val_max) and val_min_new < val_max_new:
                return val_min_new, val_max_new, True
                
        return val_min, val_max, False

    # Iterar la optimización
    for attempt in range(max_attempts):
        print(f"\n==================================================")
        print(f"Autorange: Intento {attempt + 1} de {max_attempts}")
        print(f"Límites de espesores actuales: {d_bounds_work}")
        if layer_models_work:
            print(f"Modelos de dispersión actuales:")
            for idx, lm in enumerate(layer_models_work):
                if lm:
                    name = layer_names[idx] if layer_names else f"layer_{idx}"
                    print(f"  Capa {idx} ({name}): model={lm.get('model')}, bounds={lm.get('bounds') or lm.get('f_bounds')}")
        print(f"==================================================")
        
        # Ejecutar la optimización
        if layer_models_work is not None:
            best_thicknesses, best_Is_curve, best_Ic_curve, best_params = fit_ellipsometry_torch(
                n_list=n_list, d_bounds=d_bounds_work, lams=lams, Is_exp=Is_exp, Ic_exp=Ic_exp,
                Data_R=Data_R, th_0=th_0, num_starts=num_starts, num_epochs=num_epochs, lr=lr,
                use_cuda=use_cuda, layer_models=layer_models_work, layer_names=layer_names,
                std_Is=std_Is, std_Ic=std_Ic, check_cancel_fn=check_cancel_fn,
                n_max_limits=n_max_limits
            )
        else:
            best_thicknesses, best_Is_curve, best_Ic_curve = fit_ellipsometry_torch(
                n_list=n_list, d_bounds=d_bounds_work, lams=lams, Is_exp=Is_exp, Ic_exp=Ic_exp,
                Data_R=Data_R, th_0=th_0, num_starts=num_starts, num_epochs=num_epochs, lr=lr,
                use_cuda=use_cuda, layer_models=None, layer_names=layer_names,
                std_Is=std_Is, std_Ic=std_Ic, check_cancel_fn=check_cancel_fn,
                n_max_limits=n_max_limits
            )
            best_params = None

        any_adjusted = False
        
        # 2. Verificar saturación de espesores
        is_optimizable = []
        for b in d_bounds_work:
            if isinstance(b, (tuple, list)):
                d_min, d_max = float(b[0]), float(b[1])
                is_optimizable.append(d_min != d_max)
            else:
                is_optimizable.append(False)
                
        o_idx = 0
        for idx, opt in enumerate(is_optimizable):
            if opt:
                val_opt = best_thicknesses[o_idx]
                d_min, d_max = d_bounds_work[idx]
                o_idx += 1
                
                # Ajustar rango de espesores (límite inferior absoluto es 0.0)
                new_min, new_max, adjusted = adjust_val_range(d_min, d_max, val_opt, min_abs=0.0)
                if adjusted:
                    print(f"[Autorange] Capa {idx + 2} ('{layer_names[idx + 1] if layer_names and (idx + 1) < len(layer_names) else f'layer_{idx}'}'):")
                    print(f"  El espesor se saturó cerca del límite ({val_opt:.2f} nm en rango [{d_min:.2f}, {d_max:.2f}]).")
                    print(f"  Nuevo rango propuesto: [{new_min:.2f}, {new_max:.2f}]")
                    d_bounds_work[idx] = (new_min, new_max)
                    any_adjusted = True

        # 3. Verificar saturación de parámetros de dispersión
        if layer_models_work is not None and best_params is not None:
            for idx, model in enumerate(layer_models_work):
                if model is None:
                    continue
                model_type = model.get('model', 'cauchy')
                layer_name = layer_names[idx] if layer_names else f"layer_{idx}"
                layer_params = best_params.get(idx) or best_params.get(layer_name)
                if layer_params is None:
                    continue
                
                if model_type in ['cauchy', 'cauchy_absorbent']:
                    bounds = model.get('bounds')
                    if not bounds:
                        continue
                    
                    # Mapear llaves de parámetros a índices de bounds y límites absolutos
                    if model_type == 'cauchy':
                        param_map = [('A', 1.0, None), ('B', None, None), ('C', None, None)]
                    else:
                        param_map = [('A', 1.0, None), ('B', None, None), ('C', None, None),
                                     ('D', 0.0, None), ('E', None, None), ('F', None, None)]
                    
                    for k, (key, min_abs, max_abs) in enumerate(param_map):
                        val_opt = layer_params.get(key)
                        if val_opt is None:
                            continue
                        p_min, p_max = bounds[k]
                        new_min, new_max, adjusted = adjust_val_range(p_min, p_max, val_opt, min_abs=min_abs, max_abs=max_abs)
                        if adjusted:
                            print(f"[Autorange] Capa {idx + 1} ('{layer_name}'), Parámetro {key}:")
                            print(f"  Se saturó cerca del límite ({val_opt:.4f} en rango [{p_min:.4f}, {p_max:.4f}]).")
                            print(f"  Nuevo rango propuesto: [{new_min:.4f}, {new_max:.4f}]")
                            bounds[k] = (new_min, new_max)
                            any_adjusted = True
                            
                elif model_type == 'bruggeman':
                    # Verificar f_air si es optimizable
                    if 'f_bounds' in model and model['f_bounds'] is not None:
                        f_bounds = model.get('f_bounds')
                        val_opt = layer_params.get('f_air')
                        if val_opt is not None:
                            if isinstance(f_bounds[0], (list, tuple)):
                                f_min, f_max = f_bounds[0]
                                is_nested = True
                            else:
                                f_min, f_max = f_bounds
                                is_nested = False
                            
                            new_min, new_max, adjusted = adjust_val_range(f_min, f_max, val_opt, min_abs=0.0, max_abs=1.0)
                            if adjusted:
                                print(f"[Autorange] Capa {idx + 1} ('{layer_name}'), Fracción Bruggeman (f_air):")
                                print(f"  Se saturó cerca del límite ({val_opt:.4f} en rango [{f_min:.4f}, {f_max:.4f}]).")
                                print(f"  Nuevo rango propuesto: [{new_min:.4f}, {new_max:.4f}]")
                                if is_nested:
                                    model['f_bounds'] = [(new_min, new_max)]
                                else:
                                    model['f_bounds'] = (new_min, new_max)
                                any_adjusted = True

        # Si no hubo ningún ajuste, se encontró la solución óptima sin saturar
        if not any_adjusted:
            print(f"\n[Autorange] ¡Convergencia alcanzada exitosamente en el intento {attempt + 1}!")
            break

    else:
        print(f"\n[Autorange] Se alcanzó el número máximo de intentos ({max_attempts}) sin total convergencia.")

    if layer_models_work is not None:
        return best_thicknesses, best_Is_curve, best_Ic_curve, best_params
    else:
        return best_thicknesses, best_Is_curve, best_Ic_curve

    
def transform_I_to_psi_delta(Is, Ic):
    """Convierte Is, Ic a Psi (grados) y Delta (grados)."""
    psi_deg = 0.5 * np.degrees(np.arcsin(np.clip(np.sqrt(Is**2 + Ic**2), 0.0, 1.0)))
    delta_deg = np.degrees(np.arctan2(Is, Ic))
    return psi_deg, delta_deg